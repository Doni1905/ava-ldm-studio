"""
Script to generate a publication-grade, professionally styled Microsoft Word (.docx)
document containing the complete README and research documentation for AVA LDM Studio.
"""

import os
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets background fill color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding (in twips) for a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Sets subtle custom borders for the entire table."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_callout_box(doc, text, title="KEY ARCHITECTURAL BOUNDARY"):
    """Adds an aesthetic callout / alert box with a left accent border."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F0F9FF") # Light cyan/blue
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    # Left accent border only (Sky blue #0284C7, sz 24 = 3pt)
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="0284C7"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(f"{title}: ")
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(9.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(2, 132, 199)
    
    r_body = p.add_run(text)
    r_body.font.name = "Segoe UI"
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = RGBColor(15, 23, 42)
    
    # Spacer paragraph after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    """Adds a formatted monospace code block."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "0F172A") # Dark slate
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="12" w:space="0" w:color="38BDF8"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    
    lines = code_text.strip().split("\n")
    for i, line in enumerate(lines):
        r = p.add_run(line + ("\n" if i < len(lines)-1 else ""))
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(226, 232, 240)
        
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def create_styled_table(doc, headers, rows_data, col_widths=None):
    """Creates a beautifully styled data table."""
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color="E2E8F0", sz="4")
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        set_cell_background(hdr_cells[i], "0F172A") # Deep Navy
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Segoe UI"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(248, 250, 252)
            
    # Data Rows
    for r_idx, row_data in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            for run in p.runs:
                run.font.name = "Segoe UI"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 41, 59)
                
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
                
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)
    return table

def build_document():
    doc = Document()
    
    # Page setup: Standard Letter with 0.8 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    NAVY = RGBColor(15, 23, 42)      # #0F172A
    CYAN = RGBColor(2, 132, 199)     # #0284C7
    SLATE = RGBColor(100, 116, 139)  # #64748B
    TEXT = RGBColor(30, 41, 59)      # #1E293B

    # Helper functions for headings and text
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = CYAN
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = SLATE
        return p

    def add_p(text, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Segoe UI"
            r_pre.font.size = Pt(10)
            r_pre.font.bold = True
            r_pre.font.color.rgb = NAVY
        r = p.add_run(text)
        r.font.name = "Segoe UI"
        r.font.size = Pt(10)
        r.font.italic = italic
        r.font.color.rgb = TEXT
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Segoe UI"
            r_pre.font.size = Pt(9.5)
            r_pre.font.bold = True
            r_pre.font.color.rgb = NAVY
        r = p.add_run(text)
        r.font.name = "Segoe UI"
        r.font.size = Pt(9.5)
        r.font.color.rgb = TEXT
        return p

    # ================= COVER / HEADER BLOCK =================
    header_table = doc.add_table(rows=1, cols=1)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = header_table.cell(0, 0)
    set_cell_background(c, "0F172A")
    set_cell_margins(c, top=200, bottom=200, left=240, right=240)
    
    hp = c.paragraphs[0]
    hp.paragraph_format.space_before = Pt(0)
    hp.paragraph_format.space_after = Pt(2)
    hr1 = hp.add_run("FINAL YEAR PROJECT RESEARCH REPORT\n")
    hr1.font.name = "Segoe UI"
    hr1.font.size = Pt(10)
    hr1.font.bold = True
    hr1.font.color.rgb = RGBColor(56, 189, 248) # Cyan

    hr2 = hp.add_run("AVA LDM Studio: Linguistic Dialect Model for Tamil & Tanglish\n")
    hr2.font.name = "Segoe UI"
    hr2.font.size = Pt(20)
    hr2.font.bold = True
    hr2.font.color.rgb = RGBColor(248, 250, 252)

    hr3 = hp.add_run("An On-Device Zero-Cloud Linguistic Middleware Bridging Speech Recognition & Local LLMs")
    hr3.font.name = "Segoe UI"
    hr3.font.size = Pt(11)
    hr3.font.italic = True
    hr3.font.color.rgb = RGBColor(203, 213, 225)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Executive Metadata Box
    meta_headers = ["Project Dimension", "Specification / Value"]
    meta_rows = [
        ["Project Title", "AVA LDM Studio (Accentric Virtual Assistant — Linguistic Dialect Model)"],
        ["Target Domain", "Colloquial Spoken Tamil, Regional Dialects & Tamil-English Code-Mixing (Tanglish)"],
        ["Target Dialect Varieties", "Chennai (Madras Bashai), Madurai, Kongu, Nellai, Standard Tamil"],
        ["Core Pipeline", "Speech Audio -> ASR -> Linguistic Dialect Model (LDM) -> Local LLM -> Android OS"],
        ["Execution Footprint", "On-Device Local Inference (< 1 ms LDM Latency, Zero Cloud Dependencies)"],
        ["Benchmark Performance", "53.33% End-to-End Task Success (vs 16.67% Baseline, +36.66% Absolute Improvement)"],
        ["Codebase License", "MIT License (Open Source Academic Research)"]
    ]
    create_styled_table(doc, meta_headers, meta_rows, col_widths=[2.4, 4.4])

    add_callout_box(
        doc,
        "The Linguistic Dialect Model (LDM) sits strictly between Automatic Speech Recognition (ASR) "
        "and the Local Large Language Model (LLM). The LDM normalizes dialectal Tamil/Tanglish and extracts semantic "
        "intents and slots. It MUST NOT execute Android actions directly and MUST NOT behave as a free-form conversational chatbot.",
        title="CORE SYSTEM BOUNDARY"
    )

    # ================= 1. PROJECT TITLE =================
    add_h1("1. Project Title")
    add_p("AVA LDM Studio (Accentric Virtual Assistant — Linguistic Dialect Model for Regional Tamil & Tanglish Speech).")

    # ================= 2. RESEARCH OBJECTIVE =================
    add_h1("2. Research Objective & Problem Statement")
    add_p(
        "Commercial and open-source Large Language Models (LLMs) suffer from severe comprehension degradation "
        "when confronted with regional spoken dialects, discourse fillers, and code-mixed speech (Tanglish). "
        "In South Indian contexts, conversational users interweave Tamil and English fluidly, incorporating regional colloquialisms "
        "(e.g., 'dei', 'machi', 'la', 'pa', 'ayya') and polysemous informal verbs (e.g., 'podu', 'pannu')."
    )
    add_p(
        "Standard LLMs misinterpret these discourse fillers as contact names or hallucinate refusals. "
        "The objective of AVA LDM Studio is to design, implement, and benchmark a lightweight, on-device linguistic middleware layer. "
        "Positioned between Automatic Speech Recognition (ASR) and a local quantized LLM, the LDM detects language, classifies regional dialects, "
        "strips non-semantic discourse particles, normalizes verbal inflections, and produces clean structured JSON semantic representations "
        "with zero cloud dependency."
    )

    # ================= 3. ARCHITECTURE =================
    add_h1("3. System Architecture & Component Flow")
    add_p(
        "AVA enforces a strict linear separation of concerns across a modular pipeline. "
        "The system isolates speech transcription, linguistic normalization, and action reasoning into distinct components:"
    )

    arch_diagram = (
        "+-------------------------------------------------------------------------+\n"
        "|                         INPUT: Spoken Audio Signal                      |\n"
        "|                  (16 kHz Mono PCM WAV, -20 dBFS RMS Norm)               |\n"
        "+------------------------------------+------------------------------------+\n"
        "                                     |\n"
        "                                     v\n"
        "+------------------------------------+------------------------------------+\n"
        "|                            1. ASR ENGINE                                |\n"
        "|         OpenAI Whisper (Base/Small) / Wav2Vec2-XLSR Tamil Acoustic       |\n"
        "+------------------------------------+------------------------------------+\n"
        "                                     | Raw Transcript\n"
        "                                     v\n"
        "+------------------------------------+------------------------------------+\n"
        "|                 2. LINGUISTIC DIALECT MODEL (LDM)                       |\n"
        "|   [Stage 1] Language Identification (Tamil Script, Tanglish, English)   |\n"
        "|   [Stage 2] Code-Mixing Index (CMI Token Switch Ratio)                  |\n"
        "|   [Stage 3] Regional Dialect Classifier (Chennai, Kongu, Madurai, etc.)  |\n"
        "|   [Stage 4] Morphological Normalization (Verbs: pannu, vai, anuppu)     |\n"
        "|   [Stage 5] Slang & Idiomatic Expression Mapping (Discourse filtering)  |\n"
        "|   [Stage 6] Intent Classification & Semantic Slot Extraction             |\n"
        "+------------------------------------+------------------------------------+\n"
        "                                     | Validated JSON Handoff\n"
        "                                     v\n"
        "+------------------------------------+------------------------------------+\n"
        "|                          3. LOCAL LLM INTERFACE                         |\n"
        "|        Quantized GGUF / Llama.cpp / Transformers (CPU/CUDA Offline)     |\n"
        "+------------------------------------+------------------------------------+\n"
        "                                     | Executable Action Plan\n"
        "                                     v\n"
        "+------------------------------------+------------------------------------+\n"
        "|                   4. ANDROID APPLICATION OS LAYER                       |\n"
        "|            (Executes Action: Alarm, Calendar, Navigation, Call)         |\n"
        "+-------------------------------------------------------------------------+"
    )
    add_code_block(doc, arch_diagram)

    add_h2("3.1 Dual Implementation Strategy")
    add_bullet(" High-throughput batch benchmarking, acoustic fine-tuning (Wav2Vec2), evaluation metrics, and FastAPI integration.", bold_prefix="Python Research Core (`src/`):")
    add_bullet(" Pure Kotlin 2.0 on-device rule-based engine and Material3 Compose UI running at sub-millisecond latencies (< 1 ms).", bold_prefix="Native Android Prototype (`android/`):")

    # ================= 4. DATASET =================
    add_h1("4. Dataset & Dialect Distribution")
    add_p(
        "AVA focuses on 5 canonical spoken varieties across Tamil Nadu, addressing underrepresented dialectal inflections:"
    )
    add_bullet(" High frequency of English loan elision and youth slang (machi, dei, bruh, vaada).", bold_prefix="Chennai (Madras Bashai):")
    add_bullet(" Honorific relational markers, vowel elongation, and distinctive verbal syntax (thambi, aama, ennanga).", bold_prefix="Madurai (Southern Inland):")
    add_bullet(" Softened honorific verbal suffixes (-nga) and regional particles (ayya, yov, la).", bold_prefix="Kongu (Western / Coimbatore):")
    add_bullet(" Unique phrasal verb idioms (phone pottu kudu = call someone) and particles (pa, ppa).", bold_prefix="Nellai (Southern Coastal / Tirunelveli):")
    add_bullet(" Formal and semi-formal broadcast Tamil without regional slang markers.", bold_prefix="Standard Tamil:")

    add_h2("4.1 Canonical Benchmark Dataset (30 Utterances Sample)")
    ds_headers = ["ID", "Raw Speech Input", "Dialect", "Lang", "Code-Mix", "Ground Truth Normalized Text", "Intent"]
    ds_rows = [
        ["1", "Dei nalaiku assignment submit panna remind pannu", "Chennai", "Tanglish", "Yes", "Remind me to submit my assignment tomorrow.", "CREATE_REMINDER"],
        ["2", "Machi inniku evening gym poga remind pannu", "Chennai", "Tanglish", "Yes", "Remind me to go to the gym this evening.", "CREATE_REMINDER"],
        ["3", "Naalaikku kaalaila 6 maniku alarm vai", "Standard", "Tamil (rom)", "Yes", "Set an alarm for 6 in the morning tomorrow.", "SET_ALARM"],
        ["4", "Amma ku call pannu", "Standard", "Tanglish", "Yes", "Call mother.", "MAKE_CALL"],
        ["5", "Thambi ku oru message anuppu naan late ah varen nu", "Madurai", "Tanglish", "Yes", "Send a message to my brother that I will come late.", "SEND_MESSAGE"],
        ["6", "Semma song ondru podu", "Chennai", "Tamil (rom)", "No", "Play a good song.", "PLAY_MUSIC"],
        ["7", "Inniku weather eppadi iruku", "Standard", "Tanglish", "Yes", "How is the weather today?", "CHECK_WEATHER"],
        ["8", "WhatsApp open pannu da", "Chennai", "Tanglish", "Yes", "Open WhatsApp.", "OPEN_APP"],
        ["9", "Coimbatore ku vazhi kaatu", "Kongu", "Tamil (rom)", "No", "Show me directions to Coimbatore.", "NAVIGATE"],
        ["10", "Volume konjam kammi pannu", "Standard", "Tanglish", "Yes", "Reduce the volume a little.", "DEVICE_SETTING"],
        ["14", "Appa ku phone pottu kudu", "Nellai", "Tamil (rom)", "No", "Call father.", "MAKE_CALL"],
        ["15", "Ilayaraja paatu podu la", "Kongu", "Tamil (rom)", "No", "Play Ilayaraja song.", "PLAY_MUSIC"],
        ["26", "Ayya nalaiku medicine saapida remind pannunga", "Kongu", "Tanglish", "Yes", "Remind me to take medicine tomorrow.", "CREATE_REMINDER"],
        ["29", "Friend ku call pottu kudu pa", "Nellai", "Tanglish", "Yes", "Call my friend.", "MAKE_CALL"]
    ]
    create_styled_table(doc, ds_headers, ds_rows, col_widths=[0.4, 2.0, 0.7, 0.8, 0.6, 1.8, 1.1])

    # ================= 5. INSTALLATION =================
    add_h1("5. Installation & Prerequisites")
    add_p("To clone and initialize the complete research suite:")
    add_code_block(doc, "git clone https://github.com/example/ava-ldm-studio.git\ncd ava-ldm-studio")

    # ================= 6. ENVIRONMENT SETUP =================
    add_h1("6. Environment Setup")
    add_p("Initialize Python 3.10+ virtual environment and dependencies:")
    add_code_block(
        doc,
        "# Windows PowerShell\n"
        "python -m venv .venv\n"
        ".venv\\Scripts\\Activate.ps1\n"
        "pip install --upgrade pip\n"
        "pip install -r requirements.txt\n\n"
        "# Frontend Research Studio\n"
        "npm install\n\n"
        "# Android Verification (using JDK 17/21)\n"
        ".\\android\\gradlew.bat -p android testDebugUnitTest"
    )

    # ================= 7. DATASET PREPARATION =================
    add_h1("7. Dataset Preparation & Acoustic Standards")
    add_p(
        "All speech recordings are preprocessed via `src/audio/preprocessing.py` into standard 16 kHz single-channel mono 16-bit PCM WAV. "
        "Silence below -40 dBFS is trimmed and amplitude is normalized to -20 dBFS RMS."
    )
    add_code_block(doc, "python -m src.audio.preprocessing --input raw_audio/ --output data/processed/")

    # ================= 8. MODEL SETUP =================
    add_h1("8. Model Setup & Backends")
    add_bullet(" OpenAI Whisper ('base' or 'small') configured in `src/asr/whisper_engine.py`.", bold_prefix="ASR Module:")
    add_bullet(" Fine-tuned `facebook/wav2vec2-xls-r-300m` with linear classification head.", bold_prefix="Dialect Acoustic Classifier:")
    add_bullet(" Configured via `configs/llm.yaml` supporting `llama_cpp` (GGUF CPU/CUDA), `huggingface`, or `mock` for deterministic unit testing.", bold_prefix="Local LLM Engine:")

    # ================= 9. TRAINING =================
    add_h1("9. Training & Fine-Tuning Protocols")
    add_p("To train the acoustic dialect classification model on regional speech recordings:")
    add_code_block(doc, "python -m src.dialect.train --config configs/dialect.yaml")

    # ================= 10. EVALUATION =================
    add_h1("10. Evaluation Framework & Methodology")
    add_p(
        "The evaluation framework evaluates 3 architectural paradigms over identical canonical test data:"
    )
    add_bullet(" Raw speech transcripts passed directly to Local LLM without linguistic normalization.", bold_prefix="Baseline A (`ASR -> LLM`):")
    add_bullet(" Speech normalized with generic slang mapping without regional dialect conditioning.", bold_prefix="System B (`ASR -> Generic LDM -> LLM`):")
    add_bullet(" Full AVA pipeline with acoustic/lexical dialect classification, regional phrase adaptation, and structured handoff.", bold_prefix="System C (`ASR -> Dialect-Aware LDM -> LLM`):")
    add_p("Execute the evaluation runner:")
    add_code_block(doc, "python scripts/run_evaluation.py")

    # ================= 11. RUNNING INFERENCE =================
    add_h1("11. Running Inference")
    add_p("Run the unified LDM pipeline via command line on audio or text:")
    add_code_block(
        doc,
        "# From Audio:\n"
        "python scripts/run_ldm.py --audio data/samples/chennai_sample1.wav\n\n"
        "# From Text Transcript:\n"
        "python scripts/run_ldm.py --text \"Dei nalaiku assignment submit panna remind pannu\""
    )

    # ================= 12. API USAGE =================
    add_h1("12. API Usage & Integration Layer")
    add_p("Start the FastAPI REST service for Android integration:")
    add_code_block(doc, "python scripts/serve_api.py --host 127.0.0.1 --port 8000")
    add_p("Primary Endpoint: `POST /analyze`")
    add_code_block(
        doc,
        "curl -X POST \"http://127.0.0.1:8000/analyze\" \\\n"
        "     -H \"Content-Type: application/json\" \\\n"
        "     -d '{\"text\": \"Machi inniku evening gym poga remind pannu\"}'"
    )

    # ================= 13. EXAMPLE INPUT / OUTPUT =================
    add_h1("13. Example Input / Output Walkthrough")
    io_headers = ["Pipeline Stage", "Input / Process", "Output State"]
    io_rows = [
        ["1. Speech Input", "Spoken Tamil/Tanglish audio", "16kHz Mono WAV waveform"],
        ["2. ASR Transcriber", "Acoustic sequence decoding", "\"Dei nalaiku assignment submit panna remind pannu\""],
        ["3. Language Detection", "Script & token analysis", "Language: Tanglish | Code-Mixed: true"],
        ["4. Dialect Classifier", "Slang marker 'dei' identified", "Dialect: Chennai (Confidence: 0.85)"],
        ["5. Normalization", "Discourse particle 'dei' stripped", "\"nalaiku assignment submit panna remind pannu\""],
        ["6. Verb & Time Mapping", "'nalaiku' -> tomorrow, 'panna remind pannu' -> Remind", "\"Remind me to submit my assignment tomorrow.\""],
        ["7. Intent & Slots", "Deterministic slot extraction", "Intent: CREATE_REMINDER | Slots: action, time"],
        ["8. LLM Handoff", "Standardized JSON serialization", "Structured JSON payload ready for offline LLM"]
    ]
    create_styled_table(doc, io_headers, io_rows, col_widths=[1.5, 2.5, 2.8])

    add_h2("13.1 Standardized LLM Handoff Payload Schema")
    handoff_json = (
        "{\n"
        "  \"normalized_text\": \"Remind me to submit my assignment tomorrow.\",\n"
        "  \"language\": \"Tanglish\",\n"
        "  \"dialect\": \"Chennai\",\n"
        "  \"code_mix\": true,\n"
        "  \"intent\": \"CREATE_REMINDER\",\n"
        "  \"entities\": {\n"
        "    \"action\": \"submit assignment\",\n"
        "    \"time\": \"tomorrow\"\n"
        "  },\n"
        "  \"confidence\": {\n"
        "    \"language\": 0.95,\n"
        "    \"dialect\": 0.85,\n"
        "    \"intent\": 0.90,\n"
        "    \"overall\": 0.90\n"
        "  }\n"
        "}"
    )
    add_code_block(doc, handoff_json)

    # ================= 14. RESULTS =================
    add_h1("14. Empirical Results & Performance Benchmarks")
    add_p(
        "All values presented below are verified empirical measurements recorded in `results/metrics.json`:"
    )

    res_headers = ["Metric Dimension", "Baseline A (No LDM)", "System B (Generic LDM)", "System C (AVA Dialect-Aware)", "Improvement (C vs A)"]
    res_rows = [
        ["Normalization Token F1", "37.36%", "85.21%", "85.04%", "+47.68%"],
        ["Normalization BLEU-1", "33.15%", "75.93%", "75.58%", "+42.43%"],
        ["Semantic Preservation Rate", "65.79%", "97.37%", "97.37%", "+31.58%"],
        ["Intent Accuracy", "73.33%", "76.67%", "76.67%", "+3.34%"],
        ["Entity Extraction F1", "72.13%", "75.41%", "75.41%", "+3.28%"],
        ["End-to-End Task Success Rate", "16.67%", "53.33%", "53.33%", "+36.66%"]
    ]
    create_styled_table(doc, res_headers, res_rows, col_widths=[2.1, 1.2, 1.2, 1.3, 1.0])

    add_h2("14.1 Regional Dialect Classification Metrics")
    d_headers = ["Dialect Region", "Precision", "Recall", "F1 Score", "Support (Samples)"]
    d_rows = [
        ["Chennai", "100.0%", "62.5%", "76.9%", "8"],
        ["Madurai", "100.0%", "100.0%", "100.0%", "1"],
        ["Kongu", "75.0%", "100.0%", "85.7%", "3"],
        ["Nellai", "100.0%", "100.0%", "100.0%", "2"],
        ["Standard", "88.9%", "100.0%", "94.1%", "16"],
        ["OVERALL MACRO", "92.8%", "92.5%", "91.35%", "30"]
    ]
    create_styled_table(doc, d_headers, d_rows, col_widths=[1.5, 1.3, 1.3, 1.3, 1.4])

    add_h2("14.2 Latency Breakdown & Resource Footprint")
    lat_headers = ["Component", "Mean Latency", "p50 (Median)", "p95 Latency", "Execution Mode"]
    lat_rows = [
        ["ASR Module", "12.50 ms", "12.50 ms", "12.50 ms", "Offline Feature Ingestion"],
        ["LDM Pipeline", "0.25 ms", "0.21 ms", "0.37 ms", "In-Memory Rule/Lexicon"],
        ["Local LLM Handoff", "0.02 ms", "0.01 ms", "0.02 ms", "Zero-Copy JSON Payload"],
        ["TOTAL END-TO-END", "12.77 ms", "12.72 ms", "12.88 ms", "Sub-20 ms Real-Time Budget"]
    ]
    create_styled_table(doc, lat_headers, lat_rows, col_widths=[1.8, 1.3, 1.3, 1.3, 1.5])

    # ================= 15. LIMITATIONS =================
    add_h1("15. Research Limitations & Future Work")
    add_bullet(" Out-of-vocabulary regional slang not present in curated CSV lexicons passes through without morphological normalization.", bold_prefix="Finite Lexicon Coverage:")
    add_bullet(" In short commands lacking overt dialect markers (e.g., 'Amma ku call pannu'), the model defaults conservatively to Standard Tamil.", bold_prefix="Short Utterance Ambiguity:")
    add_bullet(" Irregular phonetic romanization in Tanglish can occasionally bypass strict regex patterns.", bold_prefix="Non-Standardized Tanglish Orthography:")
    add_bullet(" The LDM strictly performs linguistic disambiguation and normalization; it does not execute device operating system actions.", bold_prefix="Execution Boundary:")

    # ================= 16. CITATION =================
    add_h1("16. Citation & Academic Attribution")
    add_p("If utilizing this architecture, dataset annotations, or benchmarks in academic work, please cite:")
    bibtex_str = (
        "@misc{ava_ldm_studio_2026,\n"
        "  title={AVA LDM Studio: A Linguistic Dialect Model for Regional Tamil Dialects and Code-Mixed Speech},\n"
        "  author={AVA Research Team},\n"
        "  year={2026},\n"
        "  howpublished={\\url{https://github.com/example/ava-ldm-studio}},\n"
        "  note={Final Year Project Research Repository}\n"
        "}"
    )
    add_code_block(doc, bibtex_str)

    # ================= 17. LICENSE =================
    add_h1("17. License Information")
    add_p(
        "This project is distributed under the MIT License. "
        "See LICENSE for full terms and conditions."
    )

    # ================= APPENDIX =================
    add_h1("Appendix: Detailed Technical Documentation Index")
    add_p("The `docs/` directory contains dedicated technical specifications:")
    app_headers = ["Document File", "Topic & Technical Scope"]
    app_rows = [
        ["docs/ARCHITECTURE.md", "System design, 3-tier boundary, and Python vs Android Kotlin contracts."],
        ["docs/DATASET.md", "Corpus composition, dialect distribution, and 30-sample benchmark metadata."],
        ["docs/LDM.md", "Detailed 6-stage linguistic normalization and morphological rule engine."],
        ["docs/ASR.md", "Speech recognition models (Whisper/Wav2Vec2) and audio preprocessing."],
        ["docs/DIALECT_MODEL.md", "Acoustic & lexical dialect classifiers and 5x5 confusion matrix analysis."],
        ["docs/LOCAL_LLM.md", "Modular local LLM adapters (llama_cpp, huggingface) and structured prompts."],
        ["docs/EVALUATION.md", "Comparative evaluation methodology across Baselines A, B, and C."],
        ["docs/EXPERIMENTS.md", "Ablation experiments, latency profiling, and qualitative failure case studies."],
        ["docs/REPRODUCIBILITY.md", "Step-by-step instructions to replicate all benchmarks and builds."],
        ["docs/LIMITATIONS.md", "Transparent disclosure of scientific boundaries and ethical privacy guarantees."]
    ]
    create_styled_table(doc, app_headers, app_rows, col_widths=[2.4, 4.4])

    output_filename = "AVA_LDM_Studio_Research_Documentation.docx"
    doc.save(output_filename)
    print(f"Successfully generated {output_filename}")

    # Copy to artifacts directory as well
    artifact_dir = r"C:\Users\HARISH KUMAR V\.gemini\antigravity\brain\2ce4d7e3-61f0-45c5-9da9-ac3ab67e5a44"
    if os.path.exists(artifact_dir):
        dest = os.path.join(artifact_dir, output_filename)
        shutil.copy(output_filename, dest)
        print(f"Copied to artifact directory: {dest}")

if __name__ == "__main__":
    build_document()
